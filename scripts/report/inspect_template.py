from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document


def points(value):
    return round(value.pt, 2) if value is not None else None


def inspect(path: Path) -> dict:
    document = Document(path)
    paragraphs = []
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text:
            paragraphs.append(
                {
                    "index": index,
                    "style": paragraph.style.name if paragraph.style else None,
                    "text": text,
                }
            )

    tables = []
    for index, table in enumerate(document.tables):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        tables.append({"index": index, "style": table.style.name if table.style else None, "rows": rows})

    styles = []
    for style in document.styles:
        if style.type == 1:
            styles.append(
                {
                    "name": style.name,
                    "font": style.font.name,
                    "size_pt": points(style.font.size),
                    "bold": style.font.bold,
                    "italic": style.font.italic,
                }
            )

    sections = []
    for index, section in enumerate(document.sections):
        sections.append(
            {
                "index": index,
                "width_in": round(section.page_width.inches, 3),
                "height_in": round(section.page_height.inches, 3),
                "top_in": round(section.top_margin.inches, 3),
                "right_in": round(section.right_margin.inches, 3),
                "bottom_in": round(section.bottom_margin.inches, 3),
                "left_in": round(section.left_margin.inches, 3),
                "header_in": round(section.header_distance.inches, 3),
                "footer_in": round(section.footer_distance.inches, 3),
                "header": [paragraph.text for paragraph in section.header.paragraphs if paragraph.text.strip()],
                "footer": [paragraph.text for paragraph in section.footer.paragraphs if paragraph.text.strip()],
            }
        )

    properties = document.core_properties
    return {
        "path": str(path),
        "paragraph_count": len(document.paragraphs),
        "nonempty_paragraphs": paragraphs,
        "table_count": len(document.tables),
        "tables": tables,
        "inline_shape_count": len(document.inline_shapes),
        "sections": sections,
        "paragraph_styles": styles,
        "properties": {
            "title": properties.title,
            "subject": properties.subject,
            "author": properties.author,
            "last_modified_by": properties.last_modified_by,
            "keywords": properties.keywords,
        },
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = inspect(args.input)
    rendered = json.dumps(result, indent=2, ensure_ascii=False)
    if args.output:
        args.output.write_text(rendered, encoding="utf-8")
    else:
        print(rendered)


if __name__ == "__main__":
    main()
